# -*- coding: utf-8 -*-
"""논문 메인 표 — 행 = 모델, 열 = 데이터 x 호라이즌, MSE / MAE 두 줄.

    python tools/paper/table_main.py                 # mfg (본문)
    python tools/paper/table_main.py --group ltsf    # LTSF (일반화)

배포 가능성을 같은 표에 표시한다.  이 논문의 요점은 "정확도 순위"와 "칩에
들어가는 것 중 순위"가 다르다는 것이라, 두 순위를 한 표에서 동시에 읽혀야 한다.

    굵게        그 열(데이터, H, 지표)에서 경량 9모델 중 1위 - 배포 여부 무시
    밑줄        2위
    †           128 KB flash / 20 KB SRAM 에 안 들어감 (실측 또는 산술 배제)
    ‡           가중치는 들어가지만 C 포트가 없어 SRAM 미검증 (CycleNet)
    ★           칩에 들어가는 것 중 1위

호라이즌을 평균 내지 않는 이유: † 와 ★ 는 호라이즌마다 다르다.  평균을 내면
"DLinear 가 H>=192 에서 flash 를 넘어 떨어져 나간다"는 핵심이 사라진다.

입력
    results/{group}_raw.csv                   정확도 (merge_raw.py)
    analysis/deployment/deploy_cells.csv      칸별 배포 상태 (table_deploy.py)
출력
    analysis/{group}/T_main.xlsx              시트: main(경량 9) / appendix(전체)
    analysis/{group}/T_main.md                같은 내용 markdown
"""
import argparse
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

GROUPS = {
    "mfg": ["Steel", "MetroPT3", "WindSCADA", "N2O", "EWELD", "BlastFurnace"],
    "ltsf": ["ETTh1", "ETTh2", "ETTm1", "ETTm2", "Weather", "Exchange",
             "Electricity", "Traffic"],
}
HS = [96, 192, 336, 720]
METRICS = ["MSE", "MAE"]

# 경량 tier.  DLinear 는 H=96 에서 칩에 들어가고 우리보다 나은 칸이 있어 반드시
# 남긴다 - 빼면 "숨겼다"가 된다 (table_mfg.py 와 같은 원칙).  우리 두 변종은 맨 아래.
EDGE = ["SparseTSF", "MixLinear", "DiPE_Linear", "SparseTSF_MLP", "FITS",
        "CycleNet", "DLinear", "FEATHer_CI", "FEATHer_mix"]
HEAVY = ["PatchTST", "TimeMixer", "TQNet", "LMS_AutoTSF", "TimesNet",
         "iTransformer", "MDMLP_EIA"]
LABEL = {"SparseTSF": "SparseTSF", "MixLinear": "MixLinear", "DiPE_Linear": "DiPE-Linear",
         "SparseTSF_MLP": "SparseTSF-MLP", "FITS": "FITS", "CycleNet": "CycleNet",
         "DLinear": "DLinear", "FEATHer_CI": "FEATHer-CI (ours)",
         "FEATHer_mix": "FEATHer-mix (ours)"}

BOLD = Font(bold=True)
BOLD_UL = Font(bold=False, underline="single")
CT = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
HDR = PatternFill("solid", fgColor="E7E6E6")
OURS = PatternFill("solid", fgColor="FFF2CC")


def load(group):
    raw = pd.read_csv(f"results/{group}_raw.csv")
    acc = raw.groupby(["model", "data", "pred_len"])[METRICS].mean()
    p96 = (raw[raw.pred_len == 96].groupby(["model", "data"]).num_params.first()
           .groupby("model").agg(["min", "max"]))
    dep = pd.read_csv("analysis/deployment/deploy_cells.csv")
    dep = dep[dep.group == group]
    state = {(r.model, r.data, int(r.H)): r.state for r in dep.itertuples()}
    return acc, p96, state


def params_label(p96, m):
    if m not in p96.index:
        return ""
    a, b = int(p96.loc[m, "min"]), int(p96.loc[m, "max"])
    return f"{a:,}" if a == b else f"{a:,}-{b:,}"


def build_cells(models, datasets, acc, state, mark_deploy):
    """칸마다 (text, rank) 를 만든다.  rank 1 = 굵게, 2 = 밑줄."""
    cells = {}
    for met in METRICS:
        cols = [(d, h) for d in datasets for h in HS]
        for d, h in cols:
            vals = {m: acc.loc[(m, d, h), met] for m in models if (m, d, h) in acc.index}
            order = sorted(vals, key=vals.get)
            dep_ok = [m for m in vals if state.get((m, d, h)) == "FITS"]
            best_dep = min(dep_ok, key=vals.get) if (mark_deploy and dep_ok) else None
            for m in models:
                if m not in vals:
                    cells[(m, met, d, h)] = ("-", 0)
                    continue
                t = f"{vals[m]:.3f}"
                if mark_deploy:
                    st = state.get((m, d, h))
                    if st in ("NO", "EXCLUDED"):
                        t += "†"
                    elif st == "UNPORTED":
                        t += "‡"
                    if m == best_dep:
                        t += "★"
                rank = (order.index(m) + 1) if m in order[:2] else 0
                cells[(m, met, d, h)] = (t, rank)
        # 평균 열: 모든 데이터 x 호라이즌
        avg = {m: acc.xs(met, axis=1).loc[m].loc[datasets].mean()
               for m in models if m in acc.index.get_level_values(0)}
        order = sorted(avg, key=avg.get)
        for m in models:
            if m in avg:
                cells[(m, met, "AVG", 0)] = (f"{avg[m]:.3f}",
                                            (order.index(m) + 1) if m in order[:2] else 0)
    return cells


def write_sheet(ws, title, models, datasets, cells, p96, note):
    ws.cell(1, 1, title).font = BOLD
    # 헤더 2줄: 데이터 / 호라이즌
    fixed = ["Model", "Params (H=96)", "Metric"]
    for i, h in enumerate(fixed, 1):
        ws.merge_cells(start_row=2, start_column=i, end_row=3, end_column=i)
        c = ws.cell(2, i, h)
        c.font, c.alignment, c.fill = BOLD, CT, HDR
    col = len(fixed) + 1
    for d in datasets:
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + len(HS) - 1)
        c = ws.cell(2, col, d)
        c.font, c.alignment, c.fill = BOLD, CT, HDR
        for j, h in enumerate(HS):
            c = ws.cell(3, col + j, h)
            c.font, c.alignment, c.fill = BOLD, CT, HDR
        col += len(HS)
    ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
    c = ws.cell(2, col, "Avg")
    c.font, c.alignment, c.fill = BOLD, CT, HDR

    r = 4
    for m in models:
        ours = m.startswith("FEATHer")
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=2)
        a = ws.cell(r, 1, LABEL.get(m, m))
        a.alignment = LEFT
        b = ws.cell(r, 2, params_label(p96, m))
        b.alignment = CT
        if ours:
            a.font = BOLD
        for k, met in enumerate(METRICS):
            rr = r + k
            ws.cell(rr, 3, met).alignment = CT
            col = len(fixed) + 1
            for d in datasets:
                for h in HS:
                    t, rank = cells[(m, met, d, h)]
                    c = ws.cell(rr, col, t)
                    c.alignment = CT
                    if rank == 1:
                        c.font = BOLD
                    elif rank == 2:
                        c.font = BOLD_UL
                    col += 1
            t, rank = cells.get((m, met, "AVG", 0), ("-", 0))
            c = ws.cell(rr, col, t)
            c.alignment = CT
            if rank == 1:
                c.font = BOLD
            elif rank == 2:
                c.font = BOLD_UL
            for cc in range(1, col + 1):
                ws.cell(rr, cc).border = THIN
                if ours:
                    ws.cell(rr, cc).fill = OURS
        r += 2
    ws.cell(r + 1, 1, note).alignment = LEFT
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 7
    ws.freeze_panes = "D4"


def markdown(models, datasets, cells, p96, mark_deploy):
    hdr = ["Model", "Params", "Metric"] + [f"{d} {h}" for d in datasets for h in HS] + ["Avg"]
    out = ["| " + " | ".join(hdr) + " |", "|" + "---|" * len(hdr)]
    for m in models:
        for k, met in enumerate(METRICS):
            row = [LABEL.get(m, m) if k == 0 else "", params_label(p96, m) if k == 0 else "", met]
            for d in datasets:
                for h in HS:
                    t, rank = cells[(m, met, d, h)]
                    row.append(f"**{t}**" if rank == 1 else (f"<u>{t}</u>" if rank == 2 else t))
            t, rank = cells.get((m, met, "AVG", 0), ("-", 0))
            row.append(f"**{t}**" if rank == 1 else (f"<u>{t}</u>" if rank == 2 else t))
            out.append("| " + " | ".join(row) + " |")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--group", default="mfg", choices=list(GROUPS))
    a = ap.parse_args()
    datasets = GROUPS[a.group]
    acc, p96, state = load(a.group)
    seeds = 3

    note_edge = ("Test MSE / MAE, seq_len=96, mean of 3 seeds, best-val-epoch selection. "
                 "Bold = best, underline = 2nd among the models shown. "
                 "† does not fit a 128 KB flash / 20 KB SRAM MCU at that (dataset, horizon) "
                 "(ARM cross-compile, stack included, or weights/input buffer alone exceed the budget). "
                 "‡ weights fit but no C port, SRAM unverified. "
                 "★ most accurate among the models that fit. "
                 "Params at H=96; a range means the count depends on the channel count D.")
    note_app = ("Full lineup incl. ceiling models, which do not fit the MCU at any cell. "
                "Shown so that model size is not the exclusion criterion. "
                "Bold / underline rank all models shown.")

    os.makedirs(f"analysis/{a.group}", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "main"
    c_edge = build_cells(EDGE, datasets, acc, state, mark_deploy=True)
    write_sheet(ws, f"Table. {a.group.upper()} accuracy and deployability (lightweight tier)",
                EDGE, datasets, c_edge, p96, note_edge)
    ws2 = wb.create_sheet("appendix")
    allm = HEAVY + EDGE
    c_all = build_cells(allm, datasets, acc, state, mark_deploy=False)
    write_sheet(ws2, f"Appendix. {a.group.upper()} accuracy, full lineup",
                allm, datasets, c_all, p96, note_app)
    out = f"analysis/{a.group}/T_main.xlsx"
    wb.save(out)

    md = (f"### {a.group.upper()} — lightweight tier\n\n"
          + markdown(EDGE, datasets, c_edge, p96, True) + "\n\n" + note_edge + "\n\n"
          f"### {a.group.upper()} — full lineup (appendix)\n\n"
          + markdown(allm, datasets, c_all, p96, False) + "\n\n" + note_app + "\n")
    with open(f"analysis/{a.group}/T_main.md", "w", encoding="utf-8") as f:
        f.write(md)
    print(f"-> {out}\n-> analysis/{a.group}/T_main.md  ({len(EDGE)} models main, {len(allm)} appendix, "
          f"{len(datasets)} datasets x {len(HS)} H x {len(METRICS)} metrics, {seeds} seeds)")


if __name__ == "__main__":
    main()
