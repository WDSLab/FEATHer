# -*- coding: utf-8 -*-
"""Manufacturing main tables — accuracy and cost — from the single-source CSVs.

Reads results/mfg_raw.csv (14 models x 6 datasets x 4 horizons x 3 seeds) and
results/complexity.csv (FFT-aware MACs), and writes analysis/mfg/:

    T1_accuracy.xlsx    per-model MSE and MAE, mean+-std over seeds
                        sheets: avg (horizon-averaged + OVERALL), per_horizon
    T2_cost.xlsx        parameters, flash bytes, MACs, deployability

Parameter counts come from mfg_raw.csv's `num_params`, i.e. what was actually
constructed for each (model, dataset, horizon) — not a hard-coded reference —
because FEATHer's count moves with both D and the horizon (822 at Steel/H96,
4,522 at BlastFurnace/H720) and a single quoted number would misrepresent it.

Deployability is decided by the horizon-WORST case, since a model has to fit
the board at every horizon it claims to serve.
"""
from __future__ import annotations

import argparse
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MFG = ["Steel", "WindSCADA", "BlastFurnace", "EWELD", "MetroPT3", "N2O"]
HS = [96, 192, 336, 720]

# NUCLEO-F103RB: 128 KB flash, minus ~15 KB firmware and ~1.5 KB tables, at
# 4 bytes per float32 parameter. See deployment/EDGE_EXPERIMENT.md.
FLASH_KB = 128
FIRMWARE_KB = 15
TABLES_KB = 1.5
MAX_PARAMS = int((FLASH_KB - FIRMWARE_KB - TABLES_KB) * 1024 / 4)

# Registry key -> printed name. "SparseTSF" alone is ambiguous once both of its
# heads are in the lineup — the linear one is the paper's 41-parameter headline
# model, the MLP one lands at FEATHer's own size — and conflating them would
# make "FEATHer beats SparseTSF" mean two different things.
DISPLAY = {
    "SparseTSF": "SparseTSF-Linear",
    "SparseTSF_MLP": "SparseTSF-MLP",
    "DiPE_Linear": "DiPE-Linear",
    "LMS_AutoTSF": "LMS-AutoTSF",
    "MDMLP_EIA": "MDMLP-EIA",
    "FEATHer": "FEATHer (ours)",
}

BEST = PatternFill("solid", fgColor="C6EFCE")
SECOND = PatternFill("solid", fgColor="FFEB9C")
HDR = PatternFill("solid", fgColor="DDDDDD")
OURS = PatternFill("solid", fgColor="FFF2CC")
NODEPLOY = PatternFill("solid", fgColor="F2F2F2")
BOLD = Font(bold=True)
CT = Alignment(horizontal="center", vertical="center")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR
        cell.font = BOLD
        cell.alignment = CT
        cell.border = THIN


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 22)


def _mark_ranks(ws, row, cols, values, deployable_only=None):
    """Green = best, yellow = 2nd, computed over the given column subset."""
    pairs = [(c, v) for c, v in zip(cols, values) if v == v]
    if deployable_only is not None:
        pairs = [(c, v) for c, v in pairs if deployable_only.get(c, True)]
    for rank, (c, _) in enumerate(sorted(pairs, key=lambda t: t[1])[:2]):
        ws.cell(row=row, column=c).fill = BEST if rank == 0 else SECOND


def build_accuracy(df: pd.DataFrame, deployable: dict, out: str) -> None:
    wb = Workbook()

    # ---------- sheet 1: horizon-averaged + OVERALL ----------
    ws = wb.active
    ws.title = "avg"
    ws.cell(row=1, column=1, value="Model").border = THIN
    ws.cell(row=1, column=2, value="Metric").border = THIN
    ws.cell(row=1, column=3, value="params").border = THIN
    ws.cell(row=1, column=4, value="MCU").border = THIN
    for j, ds in enumerate(MFG):
        ws.cell(row=1, column=5 + j, value=ds)
    ws.cell(row=1, column=5 + len(MFG), value="OVERALL")
    _style_header(ws, 1, 5 + len(MFG))

    order = (df.groupby(["model", "data"]).MSE.mean().unstack()[MFG]
             .mean(axis=1).sort_values().index.tolist())

    r = 2
    for m in order:
        sub = df[df.model == m]
        pmin, pmax = int(sub.num_params.min()), int(sub.num_params.max())
        for k, metric in enumerate(["MSE", "MAE"]):
            # mean over seeds per cell, then over horizons per dataset
            cell_mean = sub.groupby(["data", "pred_len"])[metric].mean()
            cell_std = sub.groupby(["data", "pred_len"])[metric].std()
            ws.cell(row=r, column=2, value=metric).border = THIN
            vals, cols = [], []
            for j, ds in enumerate(MFG):
                mu = cell_mean.loc[ds].mean()
                sd = cell_std.loc[ds].mean()
                c = ws.cell(row=r, column=5 + j, value=f"{mu:.4f}±{sd:.4f}")
                c.alignment, c.border = CT, THIN
                vals.append(mu)
                cols.append(5 + j)
            ov = cell_mean.groupby("data").mean()[MFG].mean()
            c = ws.cell(row=r, column=5 + len(MFG), value=round(ov, 4))
            c.alignment, c.border, c.font = CT, THIN, BOLD
            if k == 0:
                ws.cell(row=r, column=1, value=DISPLAY.get(m, m)).border = THIN
                ws.cell(row=r, column=3,
                        value=f"{pmin:,}" if pmin == pmax else f"{pmin:,}–{pmax:,}")
                ws.cell(row=r, column=4, value="yes" if deployable[m] else "no")
                for c_ in (1, 3, 4):
                    ws.cell(row=r, column=c_).alignment = CT
                    ws.cell(row=r, column=c_).border = THIN
                if m == "FEATHer":
                    for c_ in range(1, 6 + len(MFG)):
                        ws.cell(row=r, column=c_).fill = OURS
                elif not deployable[m]:
                    ws.cell(row=r, column=4).fill = NODEPLOY
            r += 1
        ws.merge_cells(start_row=r - 2, start_column=1, end_row=r - 1, end_column=1)
        for c_ in (3, 4):
            ws.merge_cells(start_row=r - 2, start_column=c_, end_row=r - 1, end_column=c_)

    # rank within each metric row-block, over the deployable models only
    for metric_off, metric in enumerate(["MSE", "MAE"]):
        for j in range(len(MFG) + 1):
            col = 5 + j
            entries = []
            for i, m in enumerate(order):
                row = 2 + 2 * i + metric_off
                v = ws.cell(row=row, column=col).value
                if v is None:
                    continue
                num = float(str(v).split("±")[0])
                entries.append((row, num, deployable[m]))
            for rank, (row, _, _) in enumerate(
                    sorted([e for e in entries if e[2]], key=lambda t: t[1])[:2]):
                ws.cell(row=row, column=col).fill = BEST if rank == 0 else SECOND
    _autosize(ws)

    # ---------- sheet 2: per horizon ----------
    ws2 = wb.create_sheet("per_horizon")
    ws2.cell(row=1, column=1, value="Model")
    ws2.cell(row=1, column=2, value="Metric")
    col = 3
    for ds in MFG:
        for h in HS:
            ws2.cell(row=1, column=col, value=f"{ds}\nH{h}").alignment = CT
            col += 1
    _style_header(ws2, 1, col - 1)
    r = 2
    for m in order:
        sub = df[df.model == m]
        for metric in ["MSE", "MAE"]:
            ws2.cell(row=r, column=1, value=DISPLAY.get(m, m)).border = THIN
            ws2.cell(row=r, column=2, value=metric).border = THIN
            g = sub.groupby(["data", "pred_len"])[metric].mean()
            col = 3
            for ds in MFG:
                for h in HS:
                    c = ws2.cell(row=r, column=col, value=round(float(g.loc[(ds, h)]), 4))
                    c.alignment, c.border = CT, THIN
                    col += 1
            r += 1
        ws2.merge_cells(start_row=r - 2, start_column=1, end_row=r - 1, end_column=1)
    _autosize(ws2)

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


def build_cost(df: pd.DataFrame, cx: pd.DataFrame, deployable: dict, out: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "cost"
    heads = ["Model", "params H96\n(min–max over datasets)", "params H720\n(min–max)",
             "worst-case flash\n(KB, x4B)", "fits 128KB MCU",
             "MACs H96\n(Steel, D=6)", "MACs H720\n(Steel)"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=1, column=j, value=h).alignment = CT
    _style_header(ws, 1, len(heads))

    order = (df.groupby("model").num_params.max().sort_values().index.tolist())
    steel_cx = cx[(cx.data == "Steel")].set_index(["model", "pred_len"]).macs_total

    r = 2
    for m in order:
        sub = df[df.model == m]
        p96 = sub[sub.pred_len == 96].num_params
        p720 = sub[sub.pred_len == 720].num_params
        worst = int(sub.num_params.max())
        vals = [
            m,
            f"{p96.min():,}" if p96.min() == p96.max() else f"{p96.min():,}–{p96.max():,}",
            f"{p720.min():,}" if p720.min() == p720.max() else f"{p720.min():,}–{p720.max():,}",
            round(worst * 4 / 1024, 1),
            "yes" if deployable[m] else "no",
            int(steel_cx.get((m, 96))) if (m, 96) in steel_cx.index else "—",
            int(steel_cx.get((m, 720))) if (m, 720) in steel_cx.index else "—",
        ]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment, c.border = CT, THIN
        if m == "FEATHer":
            for j in range(1, len(heads) + 1):
                ws.cell(row=r, column=j).fill = OURS
        elif not deployable[m]:
            ws.cell(row=r, column=5).fill = NODEPLOY
        r += 1

    r += 1
    ws.cell(row=r, column=1,
            value=(f"Deployable = worst-case params over all datasets and horizons "
                   f"fits {FLASH_KB} KB flash minus ~{FIRMWARE_KB} KB firmware and "
                   f"~{TABLES_KB} KB tables at 4 B/param, i.e. <= {MAX_PARAMS:,} params. "
                   f"MACs are FFT-aware (matmul + N*log2 N per transform); "
                   f"'—' = not measured (ceiling models are out of the deployable set)."))
    _autosize(ws)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="results/mfg_raw.csv")
    ap.add_argument("--complexity", default="results/complexity.csv")
    ap.add_argument("--out_dir", default="analysis/mfg")
    args = ap.parse_args()

    df = pd.read_csv(args.csv)
    df = df[df.data.isin(MFG)]
    cx = pd.read_csv(args.complexity)

    worst = df.groupby("model").num_params.max()
    deployable = {m: bool(v <= MAX_PARAMS) for m, v in worst.items()}

    n_cells = df.groupby(["model", "data", "pred_len"]).size()
    print(f"{len(df)} rows | {df.model.nunique()} models | "
          f"{df.data.nunique()} datasets | seeds/cell "
          f"{n_cells.min()}–{n_cells.max()}")
    print(f"MCU budget: <= {MAX_PARAMS:,} params | deployable: "
          f"{sorted(m for m, v in deployable.items() if v)}")

    build_accuracy(df, deployable, os.path.join(args.out_dir, "T1_accuracy.xlsx"))
    build_cost(df, cx, deployable, os.path.join(args.out_dir, "T2_cost.xlsx"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
